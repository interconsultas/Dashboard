"""Tests para etl_citas: lectura de la hoja de citas atendidas y carga en BD."""
import json
from datetime import date, datetime
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

import etl_citas as ec


ENCABEZADO_CLIENTE = [
    "PROFESIONAL", "ID MD", "202601", "202602", "202603",
    "PROMEDIO ", "NÚMERO DE USUARIOS ASIGNADOS ",
]


def _libro(filas, titulo="Tabla_CitasConfirmadas", otras_hojas=None):
    """Libro en memoria: `otras_hojas` es una lista de (titulo, filas) previas."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, contenido in otras_hojas or []:
        hoja = wb.create_sheet(nombre)
        for fila in contenido:
            hoja.append(fila)
    ws = wb.create_sheet(titulo)
    for fila in filas:
        ws.append(fila)
    return wb


def _hoja(filas):
    return _libro(filas)["Tabla_CitasConfirmadas"]


def _filas_cliente():
    return [
        ENCABEZADO_CLIENTE,
        ["ABDIEL SOLARTE", 75085184, 335, 360, None, 347.5, 1252],
        ["ALEJANDRA GAÑAN", 24334178, 287, None, None, 287, "NA"],
    ]


def _conn_mock(medicos):
    """`medicos`: filas (identificacion, nombre, programa_meta) del catálogo."""
    conn = MagicMock()
    cur = conn.cursor.return_value
    cur.fetchall.return_value = medicos
    return conn, cur


def _inserts(cur):
    return [
        c for c in cur.execute.call_args_list
        if c.args[0].strip().upper().startswith("INSERT INTO CITAS_ATENDIDAS")
    ]


def _guardar(tmp_path, wb, nombre="citas.xlsx"):
    ruta = tmp_path / nombre
    wb.save(ruta)
    return ruta


class TestParsearPeriodo:

    @pytest.mark.parametrize("valor, esperado", [
        (202601, 202601),
        ("202601", 202601),
        (" 202612 ", 202612),
        (202601.0, 202601),
        ("ene-26", 202601),
        ("ENE 2026", 202601),
        ("enero-2026", 202601),
        ("Feb-26", 202602),
        ("sept-26", 202609),
        ("sep.-26", 202609),
        ("dic/2025", 202512),
        ("DICIEMBRE 25", 202512),
        (datetime(2026, 3, 1), 202603),
        (date(2026, 4, 15), 202604),
    ])
    def test_reconoce_las_variantes_de_encabezado(self, valor, esperado):
        assert ec.parsear_periodo(valor) == esperado

    @pytest.mark.parametrize("valor", [
        None, "", "PROMEDIO ", "NÚMERO DE USUARIOS ASIGNADOS ", "PROFESIONAL", "ID MD",
        202613, 202600, "202613", 20261, "2026", 1252, 302.83, True,
        "xyz-26", "ene-226", "ene", "26-ene", 199912, 210101,
    ])
    def test_retorna_none_si_no_es_un_periodo(self, valor):
        assert ec.parsear_periodo(valor) is None


class TestParsearCedula:

    @pytest.mark.parametrize("valor, esperado", [
        (75085184, 75085184),
        (75085184.0, 75085184),
        ("75085184", 75085184),
        (" 24334178 ", 24334178),
    ])
    def test_acepta_enteros_positivos(self, valor, esperado):
        assert ec.parsear_cedula(valor) == esperado

    @pytest.mark.parametrize("valor", [
        None, "", "  ", "TOTAL", "12a", 0, -5, 12.5, True, "75.085.184", "1" * 16,
    ])
    def test_retorna_none_si_esta_vacia_o_no_es_numerica(self, valor):
        assert ec.parsear_cedula(valor) is None


class TestParsearCantidad:

    @pytest.mark.parametrize("valor, esperado", [
        (335, 335), (0, 0), (360.0, 360), ("98", 98), (" 12 ", 12),
    ])
    def test_acepta_enteros_mayores_o_iguales_que_cero(self, valor, esperado):
        assert ec.parsear_cantidad(valor) == esperado

    @pytest.mark.parametrize("valor", [None, "", "   "])
    def test_la_celda_vacia_no_es_un_registro(self, valor):
        assert ec.parsear_cantidad(valor) is None

    @pytest.mark.parametrize("valor", [
        -1, 12.5, "NA", "12,5", "1.200", "-3", True, 2147483648, datetime(2026, 1, 1),
    ])
    def test_rechaza_valores_que_no_son_enteros_validos(self, valor):
        with pytest.raises(ValueError):
            ec.parsear_cantidad(valor)


class TestDetectarEncabezado:

    def test_detecta_el_formato_del_cliente_e_ignora_las_demas_columnas(self):
        enc = ec.detectar_encabezado(_hoja(_filas_cliente()))

        assert enc.fila == 1
        assert enc.col_id == 1
        assert enc.col_nombre == 0
        assert enc.periodos == {2: 202601, 3: 202602, 4: 202603}

    @pytest.mark.parametrize("titulo", [
        "ID MD", "id md", " Id  Md ", "CEDULA", "CÉDULA", "Cédula",
        "DOCUMENTO", "IDENTIFICACION", "IDENTIFICACIÓN",
    ])
    def test_reconoce_las_variantes_de_la_columna_de_cedula(self, titulo):
        enc = ec.detectar_encabezado(_hoja([[titulo, "ene-26"], [123, 5]]))

        assert enc.col_id == 0
        assert enc.col_nombre is None
        assert enc.periodos == {1: 202601}

    def test_acepta_periodos_como_texto_de_mes_y_como_fecha(self):
        enc = ec.detectar_encabezado(_hoja([
            ["Profesional", "Cédula", "ene-26", datetime(2026, 2, 1), "MARZO 2026", "Total"],
            ["A", 1, 1, 2, 3, 6],
        ]))

        assert enc.periodos == {2: 202601, 3: 202602, 4: 202603}

    def test_encuentra_el_encabezado_debajo_de_filas_de_titulo(self):
        enc = ec.detectar_encabezado(_hoja([
            ["CITAS CONFIRMADAS 2026"],
            [],
            ["PROFESIONAL", "ID MD", 202601],
            ["A", 1, 10],
        ]))

        assert enc.fila == 3

    def test_no_busca_el_encabezado_mas_alla_de_la_fila_10(self):
        filas = [["relleno"] for _ in range(10)] + [["ID MD", 202601]]

        with pytest.raises(ec.ErrorArchivo):
            ec.detectar_encabezado(_hoja(filas))

    def test_falla_con_mensaje_claro_si_no_hay_columna_de_cedula(self):
        with pytest.raises(ec.ErrorArchivo, match="cédula"):
            ec.detectar_encabezado(_hoja([["PROFESIONAL", 202601], ["A", 5]]))

    def test_falla_con_mensaje_claro_si_no_hay_columnas_de_periodo(self):
        with pytest.raises(ec.ErrorArchivo, match="periodo"):
            ec.detectar_encabezado(_hoja([["PROFESIONAL", "ID MD", "PROMEDIO"], ["A", 1, 5]]))

    def test_falla_si_un_periodo_aparece_en_dos_columnas(self):
        with pytest.raises(ec.ErrorArchivo, match="202601"):
            ec.detectar_encabezado(_hoja([["ID MD", "202601", "ene-26"], [1, 5, 6]]))


class TestSeleccionarHoja:

    def test_prefiere_la_hoja_tabla_citas_confirmadas(self):
        wb = _libro(
            _filas_cliente(),
            otras_hojas=[("Otra", [["CEDULA", 202601], [1, 2]])],
        )

        ws, _ = ec.seleccionar_hoja(wb)

        assert ws.title == "Tabla_CitasConfirmadas"

    def test_el_nombre_de_la_hoja_no_distingue_mayusculas_ni_espacios(self):
        wb = _libro(
            _filas_cliente(),
            titulo=" tabla_citasconfirmadas ",
            otras_hojas=[("Otra", [["CEDULA", 202601], [1, 2]])],
        )

        ws, _ = ec.seleccionar_hoja(wb)

        assert ws.title.strip() == "tabla_citasconfirmadas"

    def test_sin_esa_hoja_usa_la_primera_cuyo_encabezado_califica(self):
        wb = _libro(
            _filas_cliente(),
            titulo="Citas",
            otras_hojas=[("Portada", [["Informe de citas"], ["2026"]])],
        )

        ws, enc = ec.seleccionar_hoja(wb)

        assert ws.title == "Citas"
        assert enc.col_id == 1

    def test_falla_si_la_hoja_esperada_no_tiene_el_formato(self):
        wb = _libro(
            [["PROFESIONAL", "PROMEDIO"]],
            otras_hojas=[("Otra", [["CEDULA", 202601], [1, 2]])],
        )

        with pytest.raises(ec.ErrorArchivo):
            ec.seleccionar_hoja(wb)

    def test_falla_si_ninguna_hoja_califica(self):
        wb = _libro([["a", "b"]], titulo="Hoja1", otras_hojas=[("Hoja0", [["x"]])])

        with pytest.raises(ec.ErrorArchivo, match="cédula"):
            ec.seleccionar_hoja(wb)

    def test_si_alguna_hoja_tiene_cedula_el_error_habla_de_los_periodos(self):
        wb = _libro([["ID MD", "PROMEDIO"]], titulo="Hoja1", otras_hojas=[("Hoja0", [["x"]])])

        with pytest.raises(ec.ErrorArchivo, match="periodo"):
            ec.seleccionar_hoja(wb)


class TestLeerRegistros:

    def _leer(self, filas):
        ws = _hoja(filas)
        return ec.leer_registros(ws, ec.detectar_encabezado(ws))

    def test_genera_un_registro_por_celda_con_valor_y_omite_las_vacias(self):
        lectura = self._leer(_filas_cliente())

        assert lectura.registros == [
            (75085184, 202601, 335),
            (75085184, 202602, 360),
            (24334178, 202601, 287),
        ]
        assert lectura.nombres == {75085184: "ABDIEL SOLARTE", 24334178: "ALEJANDRA GAÑAN"}
        assert lectura.advertencias == []

    def test_conserva_el_cero_como_registro(self):
        lectura = self._leer([["ID MD", 202601], [1, 0]])

        assert lectura.registros == [(1, 202601, 0)]

    def test_omite_las_filas_sin_cedula_o_con_cedula_no_numerica(self):
        lectura = self._leer([
            ["PROFESIONAL", "ID MD", 202601],
            ["SIN CEDULA", None, 10],
            ["TOTAL", "TOTAL", 99],
            ["VALIDO", 5, 7],
        ])

        assert lectura.registros == [(5, 202601, 7)]

    def test_reporta_y_omite_los_valores_no_enteros_o_negativos(self):
        lectura = self._leer([
            ["PROFESIONAL", "ID MD", 202601, 202602, 202603],
            ["A", 11, 12.5, 40, -3],
            ["B", 22, "NA", 8, None],
        ])

        assert lectura.registros == [(11, 202602, 40), (22, 202602, 8)]
        assert len(lectura.advertencias) == 3
        assert "Fila 2" in lectura.advertencias[0]
        assert "11" in lectura.advertencias[0]
        assert "202601" in lectura.advertencias[0]
        assert "12.5" in lectura.advertencias[0]
        assert "Fila 3" in lectura.advertencias[2]
        assert "NA" in lectura.advertencias[2]

    def test_una_cedula_repetida_conserva_la_primera_fila_y_reporta_la_otra(self):
        lectura = self._leer([
            ["PROFESIONAL", "ID MD", 202601],
            ["A", 11, 10],
            ["A BIS", 11, 99],
        ])

        assert lectura.registros == [(11, 202601, 10)]
        assert lectura.nombres == {11: "A"}
        assert len(lectura.advertencias) == 1
        assert "Fila 3" in lectura.advertencias[0]
        assert "repetida" in lectura.advertencias[0]

    def test_no_registra_al_profesional_que_no_tiene_ningun_valor(self):
        lectura = self._leer([["PROFESIONAL", "ID MD", 202601], ["A", 11, None]])

        assert lectura.registros == []
        assert lectura.nombres == {}

    def test_sin_columna_de_nombre_los_nombres_quedan_en_none(self):
        lectura = self._leer([["CEDULA", "ene-26"], [11, 4]])

        assert lectura.nombres == {11: None}


class TestLimitarAdvertencias:

    def test_no_modifica_una_lista_corta(self):
        assert ec.limitar_advertencias(["a", "b"]) == ["a", "b"]

    def test_limita_a_50_e_indica_cuantas_no_se_muestran(self):
        limitadas = ec.limitar_advertencias([f"adv {i}" for i in range(120)])

        assert len(limitadas) == 50
        assert limitadas[:49] == [f"adv {i}" for i in range(49)]
        assert "71" in limitadas[49]


class TestCargarCitas:

    MEDICOS = [
        (75085184, "SOLARTE MEJIA ABDIEL MILOVAN", "RCV"),
        (24334178, "GAÑAN LUQUE ALEJANDRA", None),
    ]

    def _cargar(self, tmp_path, filas=None, medicos=None):
        ruta = _guardar(tmp_path, _libro(filas or _filas_cliente()))
        conn, cur = _conn_mock(self.MEDICOS if medicos is None else medicos)
        with patch.object(ec, "get_db_conn", return_value=conn):
            resultado = ec.cargar_citas(ruta)
        return resultado, conn, cur

    def test_imprime_una_sola_linea_result_json_al_final(self, tmp_path, capsys):
        resultado, _, _ = self._cargar(tmp_path)

        salida = capsys.readouterr().out
        lineas_json = [l for l in salida.splitlines() if l.startswith("RESULT_JSON:")]
        assert len(lineas_json) == 1
        assert salida.strip().splitlines()[-1] == lineas_json[0]
        assert json.loads(lineas_json[0][len("RESULT_JSON:"):]) == resultado

    def test_resultado_tiene_las_claves_y_los_totales_esperados(self, tmp_path):
        resultado, _, _ = self._cargar(tmp_path)

        assert resultado == {
            "registros": 3,
            "periodos": [202601, 202602],
            "profesionales": 2,
            "sin_catalogo": [],
            "sin_programa": 1,
            "advertencias": [],
        }

    def test_consulta_el_catalogo_con_las_cedulas_como_parametro(self, tmp_path):
        _, _, cur = self._cargar(tmp_path)

        consultas = [
            c for c in cur.execute.call_args_list
            if c.args[0].strip().upper().startswith("SELECT")
        ]
        assert len(consultas) == 1
        sql, params = consultas[0].args
        assert "FROM medicos" in sql
        assert "75085184" not in sql
        assert params == ([75085184, 24334178],)

    def test_el_upsert_toma_el_nombre_y_el_programa_del_catalogo(self, tmp_path):
        _, _, cur = self._cargar(tmp_path)

        inserts = _inserts(cur)
        assert [c.args[1] for c in inserts] == [
            (75085184, 202601, 335, "SOLARTE MEJIA ABDIEL MILOVAN", "RCV"),
            (75085184, 202602, 360, "SOLARTE MEJIA ABDIEL MILOVAN", "RCV"),
            (24334178, 202601, 287, "GAÑAN LUQUE ALEJANDRA", None),
        ]

    def test_el_upsert_es_parametrizado_y_no_sobrescribe_el_programa_guardado(self, tmp_path):
        _, _, cur = self._cargar(tmp_path)

        sql = " ".join(_inserts(cur)[0].args[0].split())
        assert (
            "INSERT INTO citas_atendidas (documento, periodo, cantidad_citas, nombre_medico, programa) "
            "VALUES (%s, %s, %s, %s, %s)"
        ) in sql
        assert "ON CONFLICT (documento, periodo) DO UPDATE SET" in sql
        assert "cantidad_citas = EXCLUDED.cantidad_citas" in sql
        assert "nombre_medico = EXCLUDED.nombre_medico" in sql
        assert "programa = COALESCE(citas_atendidas.programa, EXCLUDED.programa)" in sql
        assert "335" not in sql

    def test_no_borra_nada(self, tmp_path):
        _, _, cur = self._cargar(tmp_path)

        for call in cur.execute.call_args_list:
            sql = call.args[0].strip().upper()
            assert not sql.startswith("DELETE")
            assert not sql.startswith("TRUNCATE")

    def test_hace_un_solo_commit_al_final_y_cierra_la_conexion(self, tmp_path):
        _, conn, _ = self._cargar(tmp_path)

        assert conn.commit.call_count == 1
        conn.rollback.assert_not_called()
        conn.close.assert_called_once()

    def test_omite_y_reporta_las_cedulas_que_no_estan_en_el_catalogo(self, tmp_path):
        resultado, _, cur = self._cargar(tmp_path, medicos=[self.MEDICOS[0]])

        assert [c.args[1][0] for c in _inserts(cur)] == [75085184, 75085184]
        assert resultado["registros"] == 2
        assert resultado["profesionales"] == 1
        assert resultado["periodos"] == [202601, 202602]
        assert resultado["sin_programa"] == 0
        assert resultado["sin_catalogo"] == [
            {"documento": "24334178", "nombre": "ALEJANDRA GAÑAN"},
        ]

    def test_incluye_las_advertencias_de_lectura_en_el_resultado(self, tmp_path):
        filas = [ENCABEZADO_CLIENTE, ["ABDIEL SOLARTE", 75085184, "NA", 360, None, 1, 1]]

        resultado, _, _ = self._cargar(tmp_path, filas=filas)

        assert resultado["registros"] == 1
        assert len(resultado["advertencias"]) == 1
        assert "NA" in resultado["advertencias"][0]

    def test_revierte_la_transaccion_si_una_sentencia_falla(self, tmp_path):
        ruta = _guardar(tmp_path, _libro(_filas_cliente()))
        conn, cur = _conn_mock(self.MEDICOS)

        def ejecutar(sql, params=None):
            if params and params[:2] == (75085184, 202602):
                raise RuntimeError("fallo de base de datos")

        cur.execute.side_effect = ejecutar

        with patch.object(ec, "get_db_conn", return_value=conn), \
             pytest.raises(RuntimeError, match="fallo de base de datos"):
            ec.cargar_citas(ruta)

        conn.rollback.assert_called_once()
        conn.commit.assert_not_called()
        conn.close.assert_called_once()

    def test_un_archivo_sin_el_formato_falla_antes_de_conectarse(self, tmp_path, capsys):
        ruta = _guardar(tmp_path, _libro([["PROFESIONAL", "PROMEDIO"]], titulo="Hoja1"))

        with patch.object(ec, "get_db_conn") as get_conn, \
             pytest.raises(ec.ErrorArchivo):
            ec.cargar_citas(ruta)

        get_conn.assert_not_called()
        assert "RESULT_JSON:" not in capsys.readouterr().out


class TestMain:

    def test_archivo_inexistente_termina_con_codigo_1_y_mensaje_en_stderr(self, tmp_path, capsys):
        codigo = ec.main(["--archivo", str(tmp_path / "no-existe.xlsx")])

        assert codigo == 1
        assert "no encontrado" in capsys.readouterr().err

    def test_error_de_formato_termina_con_codigo_1_y_solo_el_mensaje_en_stderr(self, tmp_path, capsys):
        ruta = _guardar(tmp_path, _libro([["PROFESIONAL", 202601]], titulo="Hoja1"))

        codigo = ec.main(["--archivo", str(ruta)])

        err = capsys.readouterr().err
        assert codigo == 1
        assert "cédula" in err
        assert "Traceback" not in err

    def test_error_de_base_de_datos_termina_con_codigo_1(self, tmp_path, capsys):
        ruta = _guardar(tmp_path, _libro(_filas_cliente()))

        with patch.object(ec, "get_db_conn", side_effect=RuntimeError("sin conexión")):
            codigo = ec.main(["--archivo", str(ruta)])

        assert codigo == 1
        assert "sin conexión" in capsys.readouterr().err

    def test_carga_exitosa_termina_con_codigo_0(self, tmp_path, capsys):
        ruta = _guardar(tmp_path, _libro(_filas_cliente()))
        conn, _ = _conn_mock(TestCargarCitas.MEDICOS)

        with patch.object(ec, "get_db_conn", return_value=conn):
            codigo = ec.main(["--archivo", str(ruta)])

        assert codigo == 0
        assert "RESULT_JSON:" in capsys.readouterr().out
